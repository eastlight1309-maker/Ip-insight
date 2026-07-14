"""분석 결과 -> PowerPoint(.pptx) 생성.

python-pptx 로 표지, 개요, 인사이트별 슬라이드(요약+항목+차트)를 만든다.
차트는 pptx 네이티브 차트를 사용해 편집 가능한 형태로 삽입한다.
"""

from __future__ import annotations

import io
from typing import Dict, List, Optional, Tuple

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from .llm_analyzer import InsightResult

# 브랜드 색 (진한 남색 + 포인트 청록)
NAVY = (0x1F, 0x2A, 0x44)
TEAL = (0x00, 0x8C, 0x99)
GRAY = (0x55, 0x5B, 0x66)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _blank_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _add_text(slide, left, top, width, height, text, size=18, bold=False,
              color=NAVY, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    return box


def _rgb(t: Tuple[int, int, int]):
    from pptx.dml.color import RGBColor

    return RGBColor(*t)


def _add_bar_chart(slide, categories: List[str], values: List[float], title: str,
                   left=Inches(7.0), top=Inches(1.6), width=Inches(5.8), height=Inches(4.9)):
    if not categories or not values:
        return
    chart_data = CategoryChartData()
    chart_data.categories = categories
    chart_data.add_series(title, values)
    gframe = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED, left, top, width, height, chart_data
    )
    chart = gframe.chart
    chart.has_legend = False
    chart.has_title = True
    chart.chart_title.text_frame.text = title
    for para in chart.chart_title.text_frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(12)


def _add_line_chart(slide, categories: List[str], values: List[float], title: str,
                    left=Inches(7.0), top=Inches(1.6), width=Inches(5.8), height=Inches(4.9)):
    if not categories or not values:
        return
    chart_data = CategoryChartData()
    chart_data.categories = categories
    chart_data.add_series(title, values)
    gframe = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS, left, top, width, height, chart_data
    )
    chart = gframe.chart
    chart.has_legend = False
    chart.has_title = True
    chart.chart_title.text_frame.text = title


def _title_slide(prs, title: str, subtitle: str):
    slide = _blank_slide(prs)
    _add_text(slide, Inches(0.9), Inches(2.6), Inches(11.5), Inches(1.4),
              title, size=40, bold=True, color=NAVY)
    _add_text(slide, Inches(0.9), Inches(4.0), Inches(11.5), Inches(1.0),
              subtitle, size=18, color=TEAL)


def _overview_slide(prs, overview: Dict[str, object]):
    slide = _blank_slide(prs)
    _add_text(slide, Inches(0.7), Inches(0.5), Inches(12), Inches(0.8),
              "데이터 개요", size=28, bold=True)
    yr = overview.get("year_range")
    yr_txt = f"{yr[0]} ~ {yr[1]}" if isinstance(yr, (list, tuple)) and yr else "N/A"
    lines = [
        f"• 총 특허 건수: {overview.get('total_records', 'N/A')} 건",
        f"• 출원 연도 범위: {yr_txt}",
        f"• 고유 출원인 수: {overview.get('unique_assignees', 'N/A')}",
        f"• 고유 발명자 수: {overview.get('unique_inventors', 'N/A')}",
    ]
    _add_text(slide, Inches(0.9), Inches(1.6), Inches(11.5), Inches(3.5),
              "\n".join(lines), size=20, color=GRAY)


def _chart_for(slide, stats: Dict[str, object]) -> bool:
    """stats 안에 그릴 수 있는 것이 있으면 차트 추가. 추가했으면 True."""
    if not isinstance(stats, dict):
        return False

    yc = stats.get("year_counts")
    if isinstance(yc, dict) and yc:
        cats = [str(k) for k in yc.keys()]
        vals = [float(v) for v in yc.values()]
        _add_line_chart(slide, cats, vals, "연도별 출원 건수")
        return True

    for key, title in (
        ("top_assignees", "상위 출원인"),
        ("top_ipc", "상위 IPC"),
        ("top_dwpi_class", "상위 DWPI Class"),
        ("top_countries", "국가별 분포"),
        ("top_inventors", "상위 발명자"),
        ("top_cpc", "상위 CPC"),
    ):
        pairs = stats.get(key)
        if isinstance(pairs, list) and pairs:
            cats = [str(p[0])[:28] for p in pairs][::-1]
            vals = [float(p[1]) for p in pairs][::-1]
            _add_bar_chart(slide, cats, vals, title)
            return True
    return False


def _insight_slide(prs, result: InsightResult):
    slide = _blank_slide(prs)
    _add_text(slide, Inches(0.7), Inches(0.4), Inches(12), Inches(0.7),
              result.name, size=26, bold=True)

    has_chart = _chart_for(slide, result.stats)
    body_width = Inches(6.0) if has_chart else Inches(12.0)

    if result.error:
        _add_text(slide, Inches(0.7), Inches(1.4), body_width, Inches(5.5),
                  f"분석 오류: {result.error}", size=14, color=(0xB0, 0x00, 0x20))
        return

    # 요약
    top = Inches(1.3)
    if result.summary:
        box = _add_text(slide, Inches(0.7), top, body_width, Inches(1.2),
                        result.summary, size=14, color=TEAL, bold=True)
        top = Inches(2.5)

    # 항목들
    box = slide.shapes.add_textbox(Inches(0.7), top, body_width, Inches(4.8))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for item in result.items:
        # 제목 문단
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        r = p.add_run()
        r.text = f"▪ {item.get('title', '')}"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = _rgb(NAVY)
        # 내용 문단
        p2 = tf.add_paragraph()
        r2 = p2.add_run()
        r2.text = item.get("content", "")
        r2.font.size = Pt(11)
        r2.font.color.rgb = _rgb(GRAY)


def build_ppt(
    results: List[InsightResult],
    overview: Optional[Dict[str, object]] = None,
    title: str = "IP 인사이트 분석 리포트",
    subtitle: str = "Derwent DWPI 기반 특허 포트폴리오 분석",
) -> bytes:
    """분석 결과로 PPT를 만들어 바이트로 반환."""
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    _title_slide(prs, title, subtitle)
    if overview:
        _overview_slide(prs, overview)

    for result in results:
        _insight_slide(prs, result)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
